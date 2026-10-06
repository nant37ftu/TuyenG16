#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Xuat toan bo noi dung dang hien tren web ra 1 file Excel de team sua.

Chay:  python tools/xuat-excel.py
Ra:    ../37FTU-G16-Noi-dung-web.xlsx  (ngoai thu muc web/ nen khong len repo cong khai)

Team sua truc tiep vao o. Cot "Ma" la khoa de nap nguoc lai, dung sua.
Khong dua so dien thoai / email / MSSV cua ai vao file nay.
"""

import argparse
import json
import os
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

GOC = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

CAM = 'FA501E'
VANG = 'FFF7D6'
XAM = '6B6B6B'
SO_HUYEN = 21

dam_trang = Font(bold=True, color='FFFFFF', size=11)
vien_mong = Border(*[Side(style='thin', color='D9D9D9')] * 4)

S_HUYEN = 'Huyện (chọn sẵn)'
GHI_CHU = 'Ghi chú của team'


# ---------------------------------------------------------------- doc du lieu

def doc_js(ten_file, ten_bien):
    """Lay object JS gan cho window.<ten_bien> roi doi thanh dict Python.

    Quet tung ky tu de khong nham: // trong URL khong phai chu thich, dau nhay
    don trong chu thich khong phai chuoi.
    """
    s = open(os.path.join(GOC, ten_file), encoding='utf-8').read()
    i = s.index('{', s.index('window.' + ten_bien))
    ra = []          # cac manh JSON se ghep lai
    muc = 0
    n = len(s)
    while i < n:
        c = s[i]
        if c == '/' and s[i + 1:i + 2] == '/':                   # chu thich 1 dong
            i = s.find('\n', i)
            if i < 0:
                break
        elif c == '/' and s[i + 1:i + 2] == '*':                 # chu thich nhieu dong
            j = s.find('*/', i + 2)
            i = n if j < 0 else j + 2
        elif c in '"\'':                                         # chuoi
            dau, i, buf = c, i + 1, []
            while i < n and s[i] != dau:
                if s[i] == '\\':
                    buf.append(s[i:i + 2])
                    i += 2
                    continue
                buf.append(s[i])
                i += 1
            i += 1
            chu = ''.join(buf).replace('\\\'', '\'').replace('\\"', '"')
            ra.append(json.dumps(chu, ensure_ascii=False))
        elif c.isalpha() or c == '_':                             # ten khoa hoac true/false/null
            j = i
            while j < n and (s[j].isalnum() or s[j] == '_'):
                j += 1
            tu = s[i:j]
            k = j
            while k < n and s[k] in ' \t\r\n':
                k += 1
            ra.append(json.dumps(tu) if s[k:k + 1] == ':' else tu)
            i = j
        else:
            if c in '}]':                                        # bo dau phay thua truoc dau dong
                while ra and ra[-1].strip() == '':
                    ra.pop()
                if ra and ra[-1] == ',':
                    ra.pop()
            ra.append(c)
            i += 1
            if c == '{':
                muc += 1
            elif c == '}':
                muc -= 1
                if muc == 0:
                    break
    return json.loads(''.join(ra))


def doc_json(ten):
    return json.load(open(os.path.join(GOC, 'data', ten), encoding='utf-8'))


def doc_json_neu_co(*duong_dan):
    """Thử lần lượt vài chỗ, không có thì trả về None (không làm hỏng cả file Excel).

    thanh-vien.json đi theo bản đồ người Nghệ: tính năng tạm ngưng nên file nằm
    trong rieng-tu/ban-do/data/, không còn trong repo công khai.
    """
    for p in duong_dan:
        p = os.path.join(GOC, p)
        if os.path.isfile(p):
            return json.load(open(p, encoding='utf-8'))
    return None


def bo_the(s):
    s = re.sub(r'<br\s*/?>', '\n', s)
    s = re.sub(r'<[^>]+>', '', s)
    s = re.sub(r'&amp;', '&', s)
    return re.sub(r'[ \t]+', ' ', s).strip()


def doc_chu_html(ten_file):
    """Lay cac dong chu bien tap duoc trong 1 trang HTML, kem so dong."""
    ra = []
    mau = [
        (r'<span class="nhan">(.*?)</span>', 'Nhãn mục'),
        (r'<h1[^>]*>(.*?)</h1>', 'Tiêu đề lớn'),
        (r'<h2[^>]*>(.*?)</h2>', 'Tiêu đề mục'),
        (r'<h3[^>]*>(.*?)</h3>', 'Tiêu đề nhỏ'),
        (r'<p class="dan-de"[^>]*>(.*?)</p>', 'Câu dẫn'),
    ]
    duong = os.path.join(GOC, ten_file)
    if not os.path.isfile(duong):
        return ra          # trang tạm ngưng (vd ban-do.html) thì bỏ qua, không lỗi
    noi_dung = open(duong, encoding='utf-8').read()
    for so, dong in enumerate(noi_dung.split('\n'), 1):
        for bieu, loai in mau:
            for m in re.finditer(bieu, dong, flags=re.S):
                chu = bo_the(m.group(1))
                if chu and '{{' not in chu:
                    ra.append((ten_file + ':' + str(so), loai, chu))
    return ra


def doc_don(ten_file='index.html'):
    """Lay tung o trong don ung tuyen: nhan, goi y, loi bao, lua chon."""
    s = open(os.path.join(GOC, ten_file), encoding='utf-8').read()
    vi_tri = [m.start() for m in re.finditer(r'data-truong="', s)]
    ra = []
    for k, dau in enumerate(vi_tri):
        het = vi_tri[k + 1] if k + 1 < len(vi_tri) else min(
            s.index('<button', dau), s.index('</form>', dau))
        khoi = s[dau:het]
        truong = re.search(r'data-truong="([^"]+)"', khoi).group(1)
        nhan = re.search(r'<label[^>]*>(.*?)</label>', khoi, flags=re.S)
        # o dong y: data-truong nam tren chinh the <label> nen khong khop mau tren
        nhan = nhan.group(1) if nhan else khoi.split('</label>')[0].split('>', 1)[-1]
        nhan = bo_the(nhan)
        bat_buoc = 'Có' if 'class="sao"' in khoi else ''
        nhan = nhan.replace('*', '').strip()
        cho = re.search(r'placeholder="([^"]*)"', khoi)
        chon = [bo_the(x) for x in re.findall(r'<option[^>]*>(.*?)</option>', khoi)]
        chon = [x for x in chon if x and not x.startswith('—')]
        goi_y = re.search(r'<p class="goi-y"[^>]*>(.*?)</p>', khoi, flags=re.S)
        loi = re.search(r'<p class="loi"[^>]*>(.*?)</p>', khoi, flags=re.S)
        ra.append([truong, nhan, bat_buoc, cho.group(1) if cho else '',
                   '\n'.join(chon), bo_the(goi_y.group(1)) if goi_y else '',
                   bo_the(loi.group(1)) if loi else '', ''])

    # chu tren nut gui va cau nhac ngay duoi nut
    nut = re.search(r'<button[^>]*id="nut-gui"[^>]*>(.*?)</button>', s, flags=re.S)
    if nut:
        ra.append(['nut_gui', bo_the(nut.group(1)), '', '', '', '', '', ''])
    sau_nut = re.search(r'</button>\s*<p class="goi-y"[^>]*>(.*?)</p>', s, flags=re.S)
    if sau_nut:
        ra.append(['chu_duoi_nut', bo_the(sau_nut.group(1)), '', '', '', '', '', ''])
    return ra


# ---------------------------------------------------------------- ve bang

def to_sheet(wb, ten, ghi_chu, cot, dong, loc=True):
    """cot = [(tieu de, do rong)], dong = list cac list o."""
    ws = wb.create_sheet(ten)
    ws.sheet_properties.tabColor = CAM

    ws.cell(1, 1, ghi_chu).font = Font(italic=True, color=XAM, size=10)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(len(cot), 1))
    ws.row_dimensions[1].height = 32
    ws.cell(1, 1).alignment = Alignment(vertical='center', wrap_text=True)

    for j, (tieu_de, w) in enumerate(cot, 1):
        o = ws.cell(2, j, tieu_de)
        o.font = dam_trang
        o.fill = PatternFill('solid', fgColor=CAM)
        o.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[2].height = 30

    for i, hang in enumerate(dong, 3):
        for j, gia_tri in enumerate(hang, 1):
            o = ws.cell(i, j, gia_tri)
            o.alignment = Alignment(vertical='top', wrap_text=True)
            o.border = vien_mong
            if i % 2 == 0:
                o.fill = PatternFill('solid', fgColor='FBFBFB')
        ws.cell(i, len(cot)).fill = PatternFill('solid', fgColor=VANG)   # cot Ghi chu

    ws.freeze_panes = 'A3'
    if loc and dong:
        ws.auto_filter.ref = 'A2:' + get_column_letter(len(cot)) + str(len(dong) + 2)
    return ws


def bang_khoa_gia_tri(wb, ten, ghi_chu, dong):
    return to_sheet(wb, ten, ghi_chu, [
        ('Mã (đừng sửa)', 26), ('Chỗ hiện trên web', 30),
        ('Nội dung — sửa trực tiếp vào ô này', 78), (GHI_CHU, 26),
    ], [[a, b, c, ''] for a, b, c in dong], loc=False)


# ---------------------------------------------------------------- cac trang

def trang_huong_dan(wb, tv, mon, ten_sheet, ban_do_chay):
    ws = wb.create_sheet('Hướng dẫn', 0)
    ws.sheet_properties.tabColor = '333333'
    ws.column_dimensions['A'].width = 32
    ws.column_dimensions['B'].width = 106
    ws.sheet_view.showGridLines = False
    hang = [0]

    def dong(a, b='', to=False, cao=None):
        hang[0] += 1
        r = hang[0]
        ws.cell(r, 1, a).font = Font(bold=True, size=13 if to else 11,
                                     color=CAM if to else '222222')
        ws.cell(r, 2, b).alignment = Alignment(vertical='top', wrap_text=True)
        if cao:
            ws.row_dimensions[r].height = cao

    so_nguoi = len(tv.get('nguoi', []))
    dong('NỘI DUNG WEB TUYỂN GEN 16 — 37FTU',
         'Toàn bộ chữ đang hiện trên trang, bày ra thành bảng để cả team đọc và sửa.', to=True)
    dong('Trang web', 'https://nant37ftu.github.io/TuyenG16/')
    dong('Ngày xuất bản này', tv.get('cap_nhat', ''))
    dong('')
    dong('SỬA THẾ NÀO', 'Sửa trực tiếp vào ô chữ. Không cần tô màu, không cần ghi lại câu cũ — '
                        'Trí sẽ so với bản gốc để biết chỗ nào đổi.', to=True, cao=32)
    dong('Cột "Mã"', 'Đừng sửa, đừng xoá dòng. Đó là khoá để đưa chữ trở lại đúng chỗ trên web.')
    dong('Cột "Ghi chú của team"', 'Ô vàng cuối mỗi dòng: viết ý kiến, câu hỏi, hoặc "bỏ dòng này".')
    dong('Muốn thêm dòng mới', 'Thêm ở cuối sheet, cột Mã để trống, ghi chú "thêm mới".')
    dong('Xuống dòng trong 1 ô', 'Alt + Enter. Mỗi lần xuống dòng là một ý riêng (ví dụ cột Thành tích).')
    dong('')
    if so_nguoi and ban_do_chay:
        dong('ƯU TIÊN LÀM TRƯỚC',
             'Sheet "Bản đồ - Người" có ' + str(so_nguoi) + ' người nhưng cột Quê đang trống hết '
             '(' + str(tv.get('da_biet_que', 0)) + '/' + str(so_nguoi) + ' người có quê), nên bản đồ trên web '
             'đang trắng. Bấm vào ô ở cột Quê sẽ có danh sách ' + str(SO_HUYEN) + ' huyện để chọn.',
             to=True, cao=48)
    else:
        dong('BẢN ĐỒ NGƯỜI NGHỆ',
             'Tính năng đang tạm ngưng, trang bản đồ không còn trên web. Sheet "Bản đồ - Người" '
             'chỉ để team chuẩn bị sẵn dữ liệu (nhất là cột Quê) cho lúc bật lại; '
             'file gốc giữ offline trong rieng-tu/ban-do/.', to=True, cao=46)
    dong('Đừng đưa vào file này',
         'Số điện thoại, email, ngày sinh, mã sinh viên của bất kỳ ai. Web là trang công khai, '
         'chỉ nên nêu tên + chức vụ của người đã đồng ý.', cao=32)
    dong('')
    dong('CÁC SHEET', '', to=True)
    y_nghia = {
        'Trang chủ': 'Tên thế hệ, tiêu đề lớn, đoạn giới thiệu',
        'Chữ trên trang': 'Tiêu đề, nhãn mục, câu dẫn của 5 trang',
        'Số liệu': '4 con số lớn ở đầu trang chủ',
        'Giá trị cốt lõi': '4 giá trị của CLB',
        'Quy tắc 4T': 'Tôn trọng / Trách nhiệm / Tự giác / Tin tưởng',
        'Ba ban': 'Mô tả 3 ban, biệt danh, học được gì, hợp với ai',
        'Quyền lợi': '4 thứ thành viên nhận được',
        'Vòng tuyển': '5 vòng + mốc thời gian',
        'Hỏi đáp': 'Câu hỏi thường gặp ở cuối trang chủ',
        'Đơn ứng tuyển': 'Từng câu hỏi trong đơn, câu gợi ý và câu báo lỗi',
        'Kênh biết đến': 'Lựa chọn trong đơn: biết về CLB qua đâu',
        'Cấu hình': 'Hạn nộp đơn, fanpage, email, hotline',
        'Bản đồ - Người': str(so_nguoi) + ' người hiện trên bản đồ — quan trọng nhất',
        S_HUYEN: 'Danh sách ' + str(SO_HUYEN) + ' huyện/thị/thành, chỉ để tham chiếu',
        'Trắc nghiệm': '10 câu "Em hợp ban nào" + điểm từng đáp án',
        'Món ăn': str(len(mon.get('mon', []))) + ' món của trang "Hôm nay ăn chi"',
        'Lý do ăn': 'Câu trả lời vui khi quay ra món',
        'Từ tiếng Nghệ': 'Từ vựng trò "Giọng Nghệ tốc độ"',
    }
    for t in ten_sheet:
        dong('  ' + t, y_nghia.get(t, ''))
    dong('')
    dong('GỬI LẠI', 'Sửa xong gửi lại đúng file này cho Trí. Trí đọc file rồi đổi thẳng lên web, '
                    'không phải gõ tay lại.', cao=30)
    return ws


def trang_ban_do(wb, tv):
    nguoi = tv.get('nguoi', [])
    cot = [
        ('Mã (đừng sửa)', 22), ('Tên', 22), ('Gen', 6), ('Quê (huyện)', 18),
        ('Nhóm', 16), ('Chức vụ', 30), ('Ban', 14), ('Nhiệm kỳ', 22),
        ('Thành tích (Alt+Enter mỗi ý)', 34), ('Lời nhắn', 34),
        ('Link Facebook', 26), ('Hành trình (chỉ để xem)', 34),
        ('Gỡ tên khỏi bản đồ? (ghi x)', 14), (GHI_CHU, 24),
    ]
    dong = []
    for p in nguoi:
        hanh_trinh = '\n'.join(
            (h.get('nk', '') + ': ' + (h.get('chuc_vu') or '') +
             (' (' + h['ban'] + ')' if h.get('ban') else ''))
            for h in (p.get('hanh_trinh') or []))
        dong.append([
            p.get('ma', ''), p.get('ten', ''), p.get('gen', ''), p.get('que', ''),
            p.get('nhom', ''), p.get('chuc_vu', ''), p.get('ban', ''),
            ', '.join(p.get('nhiem_ky') or []),
            '\n'.join(p.get('thanh_tich') or []), p.get('loi_nhan', ''),
            p.get('lien_he', ''), hanh_trinh, '', '',
        ])
    ws = to_sheet(wb, 'Bản đồ - Người',
                  'Những người hiện ra khi bấm vào một huyện trên bản đồ. Cột Quê đang trống nên bản đồ '
                  'còn trắng. Chỉ nêu tên người đã đồng ý; ai muốn rút thì ghi x ở cột "Gỡ tên khỏi bản đồ".',
                  cot, dong)
    for r in range(3, len(dong) + 3):
        ws.cell(r, 4).fill = PatternFill('solid', fgColor=VANG)
    kt = DataValidation(type='list', allow_blank=True, showDropDown=False,
                        formula1="'" + S_HUYEN + "'!$A$3:$A$" + str(SO_HUYEN + 2))
    kt.errorTitle = 'Tên huyện chưa đúng'
    kt.error = 'Chọn đúng tên trong sheet "' + S_HUYEN + '".'
    ws.add_data_validation(kt)
    kt.add('D3:D' + str(len(dong) + 2))
    return ws


def trang_trac_nghiem(wb):
    cot = [('Mã (đừng sửa)', 14), ('Câu hỏi', 46), ('Đáp án', 52),
           ('Điểm Ban Tổ chức', 12), ('Điểm Ban Truyền thông', 14),
           ('Điểm Ban Đối ngoại', 12), (GHI_CHU, 24)]
    s = open(os.path.join(GOC, 'quiz.js'), encoding='utf-8').read()
    khoi = s[s.index('var CAU_HOI = ['):]
    khoi = khoi[:khoi.index('\n  ];') + 5]
    dong = []
    cau_hien, so_cau, so_dap = '', 0, 0
    for d in khoi.split('\n'):
        m = re.search(r"\{\s*hoi:\s*'((?:[^'\\]|\\.)*)'", d)
        if m:
            so_cau, so_dap = so_cau + 1, 0
            cau_hien = m.group(1).replace("\\'", "'")
            continue
        m = re.search(r"\{\s*chu:\s*'((?:[^'\\]|\\.)*)',\s*d:\s*\[([0-9, ]+)\]", d)
        if m:
            so_dap += 1
            diem = [int(x) for x in m.group(2).split(',')]
            dong.append(['cau%d.%d' % (so_cau, so_dap),
                         cau_hien if so_dap == 1 else '',
                         m.group(1).replace("\\'", "'"), diem[0], diem[1], diem[2], ''])
    return to_sheet(wb, 'Trắc nghiệm',
                    'Trang "Em hợp ban nào?". Câu hỏi ghi ở dòng đầu của mỗi nhóm 3 đáp án. '
                    'Đáp án cộng điểm cho ban nào thì điền số vào cột điểm của ban đó.',
                    cot, dong, loc=False)


# ---------------------------------------------------------------- chay

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--ra', default=os.path.join(os.path.dirname(GOC), '37FTU-G16-Noi-dung-web.xlsx'))
    tuy = ap.parse_args()

    nd = doc_js('noi-dung.js', 'NOI_DUNG')
    ch = doc_js('config.js', 'CAU_HINH')
    # Bản đồ tạm ngưng: file nằm ở rieng-tu/ban-do/data/. Không có cũng chạy được.
    tv = doc_json_neu_co('data/thanh-vien.json',
                         'rieng-tu/ban-do/data/thanh-vien.json') or {}
    ban_do_chay = os.path.isfile(os.path.join(GOC, 'ban-do.html'))
    mon = doc_json('mon-an.json')
    huyen = doc_json('nghe-an.json')['units']

    wb = Workbook()
    wb.remove(wb.active)

    # --- trang chu
    bang_khoa_gia_tri(wb, 'Trang chủ',
                      'Ba dòng chữ chính ở đầu trang chủ (file noi-dung.js).',
                      [('the_he', 'Tên thế hệ, hiện khắp trang', nd['the_he']),
                       ('tieu_de', 'Dòng chữ lớn nhất trang chủ', nd['tieu_de']),
                       ('phu_de', 'Đoạn giới thiệu ngay dưới tiêu đề', nd['phu_de'])])

    # --- chu trong HTML
    dong = []
    ten_trang = {'index.html': 'Trang chủ', 'ban-do.html': 'Bản đồ', 'quiz.html': 'Trắc nghiệm',
                 'an-gi.html': 'Ăn chi', 'game.html': 'Game'}
    for f, nhan in ten_trang.items():
        for ma, loai, chu in doc_chu_html(f):
            dong.append([ma, nhan, loai, chu, ''])
    to_sheet(wb, 'Chữ trên trang',
             'Tiêu đề và câu dẫn nằm trong file HTML. Sửa chữ ở cột "Nội dung"; xuống dòng trong ô = ngắt dòng trên web.',
             [('Mã (đừng sửa)', 22), ('Trang', 14), ('Loại', 14),
              ('Nội dung — sửa trực tiếp', 78), (GHI_CHU, 24)], dong)

    # --- so lieu
    to_sheet(wb, 'Số liệu', '4 con số lớn ở đầu trang chủ. Con số 429 đếm thật từ sheet thành viên.',
             [('Mã (đừng sửa)', 14), ('Con số', 14), ('Chữ dưới con số', 40), (GHI_CHU, 30)],
             [['so_lieu' + str(i + 1), x['so'], x['nhan'], ''] for i, x in enumerate(nd['so_lieu'])], loc=False)

    # --- gia tri
    to_sheet(wb, 'Giá trị cốt lõi', 'Mục "Chúng tôi là ai" trên trang chủ.',
             [('Mã (đừng sửa)', 14), ('Tên giá trị', 24), ('Mô tả', 86), (GHI_CHU, 26)],
             [['gia_tri' + str(i + 1), x['ten'], x['mo_ta'], ''] for i, x in enumerate(nd['gia_tri'])], loc=False)

    # --- 4T
    to_sheet(wb, 'Quy tắc 4T', 'Bốn chữ T ngay dưới phần giá trị cốt lõi.',
             [('Mã (đừng sửa)', 14), ('Chữ', 18), ('Ý nghĩa', 86), (GHI_CHU, 26)],
             [['4t' + str(i + 1), x['chu'], x['y'], ''] for i, x in enumerate(nd['quy_tac_4t'])], loc=False)

    # --- ba ban
    to_sheet(wb, 'Ba ban',
             'Mục "Em thuộc về đâu". Cột "Học được gì": mỗi ý một dòng (Alt+Enter). '
             'Trang trắc nghiệm lấy lại chính các mô tả này.',
             [('Mã (đừng sửa)', 14), ('Tên ban', 20), ('Biệt danh', 22), ('Tóm tắt 1 dòng', 30),
              ('Mô tả đầy đủ', 70), ('Học được gì (mỗi ý 1 dòng)', 34), ('Hợp với ai', 40), (GHI_CHU, 24)],
             [['ban' + str(i + 1), b['ten'], b.get('biet_danh', ''), b.get('tom_tat', ''), b.get('mo_ta', ''),
               '\n'.join(b.get('hoc_duoc') or []), b.get('hop_voi', ''), '']
              for i, b in enumerate(nd['cac_ban'])], loc=False)

    # --- quyen loi
    to_sheet(wb, 'Quyền lợi', 'Mục "Vào rồi thì được gì" trên trang chủ.',
             [('Mã (đừng sửa)', 14), ('Tên', 34), ('Mô tả', 80), (GHI_CHU, 26)],
             [['nhan' + str(i + 1), x['ten'], x['mo_ta'], ''] for i, x in enumerate(nd['nhan_duoc'])], loc=False)

    # --- vong tuyen
    to_sheet(wb, 'Vòng tuyển',
             'Mục "Đường vào nhà". Sửa mốc thời gian ở đây; riêng hạn nộp đơn còn phải sửa ở sheet Cấu hình.',
             [('Mã (đừng sửa)', 14), ('Tên vòng', 24), ('Thời gian', 26), ('Mô tả', 70), (GHI_CHU, 26)],
             [['vong' + str(i + 1), x['ten'], x['thoi_gian'], x['mo_ta'], ''] for i, x in enumerate(nd['vong_tuyen'])],
             loc=False)

    # --- hoi dap
    to_sheet(wb, 'Hỏi đáp', 'Phần hỏi đáp ở cuối trang chủ. Thêm câu mới cứ thêm dòng ở dưới.',
             [('Mã (đừng sửa)', 14), ('Câu hỏi', 46), ('Câu trả lời', 84), (GHI_CHU, 26)],
             [['hoi' + str(i + 1), x['hoi'], x['dap'], ''] for i, x in enumerate(nd['cau_hoi'])], loc=False)

    # --- don ung tuyen
    to_sheet(wb, 'Đơn ứng tuyển',
             'Từng ô trong đơn ứng tuyển. "Câu gợi ý" hiện mờ dưới ô, "Câu báo lỗi" hiện khi ứng viên điền sai. '
             'Các lựa chọn của Quê / Nguyện vọng / Biết qua đâu lấy từ sheet khác nên để trống ở đây.',
             [('Mã (đừng sửa)', 18), ('Câu hỏi hiện trong đơn', 40), ('Bắt buộc', 10),
              ('Chữ mờ trong ô', 26), ('Lựa chọn có sẵn', 24), ('Câu gợi ý', 40),
              ('Câu báo lỗi', 40), (GHI_CHU, 24)],
             doc_don(), loc=False)

    # --- kenh biet den
    to_sheet(wb, 'Kênh biết đến', 'Lựa chọn trong đơn: "Em biết 37FTU qua đâu?".',
             [('Mã (đừng sửa)', 14), ('Lựa chọn', 40), (GHI_CHU, 30)],
             [['kenh' + str(i + 1), x, ''] for i, x in enumerate(nd['kenh_biet_den'])], loc=False)

    # --- cau hinh
    bang_khoa_gia_tri(wb, 'Cấu hình',
                      'Mốc thời gian và liên hệ (file config.js). Khoá kỹ thuật của Supabase không đưa vào đây, '
                      'cứ để trong code.',
                      [('HAN_NOP_DON', 'Đồng hồ đếm ngược + hạn trên đơn', ch['HAN_NOP_DON']),
                       ('DANG_MO_DON', 'true = còn nhận đơn, false = ẩn form', str(ch['DANG_MO_DON']).lower()),
                       ('FANPAGE', 'Link fanpage ở chân trang', ch['FANPAGE']),
                       ('EMAIL', 'Email liên hệ ở chân trang', ch['EMAIL']),
                       ('HOTLINE', 'Số gọi cho BTC (để trống = ẩn)', ch.get('HOTLINE', '')),
                       ('LINK_FORM_DU_PHONG', 'Google Form dự phòng (để trống = dùng form trên trang)',
                        ch.get('LINK_FORM_DU_PHONG', ''))])

    # --- ban do + huyen (bo qua khi tinh nang dang tam ngung)
    if tv.get('nguoi'):
        trang_ban_do(wb, tv)
        to_sheet(wb, S_HUYEN,
                 'Danh sách ' + str(SO_HUYEN) + ' huyện/thị/thành vẽ trên bản đồ. Chỉ để tham chiếu, đừng đổi tên ở đây.',
                 [('Tên hiện trên bản đồ', 26), ('Tên đầy đủ', 30), ('Mã vùng', 20), (GHI_CHU, 26)],
                 [[h['name'], h.get('full', ''), h['id'], ''] for h in sorted(huyen, key=lambda x: x['name'])],
                 loc=False)

    # --- trac nghiem
    trang_trac_nghiem(wb)

    # --- mon an
    ten_buoi = {'sang': 'Sáng', 'trua': 'Trưa', 'chieu': 'Chiều', 'toi': 'Tối', 'dem': 'Khuya'}
    ten_kieu = {'que': 'Món quê', 'pho-thong': 'Món thường', 'vat': 'Ăn vặt'}
    to_sheet(wb, 'Món ăn',
             'Trang "Hôm nay ăn chi". Giá = nghìn đồng (chỉ để lọc túi tiền). '
             'Kiểu chỉ được ghi: Món quê / Món thường / Ăn vặt. Buổi: Sáng, Trưa, Chiều, Tối, Khuya (cách nhau dấu phẩy).',
             [('Mã (đừng sửa)', 12), ('Icon', 8), ('Tên món', 26), ('Kiểu', 14),
              ('Giá (nghìn)', 10), ('Ăn buổi nào', 24), ('Mô tả', 62),
              ('Chỗ ăn quen của CLB', 26), (GHI_CHU, 22)],
             [['mon' + str(i + 1), m.get('icon', ''), m['ten'], ten_kieu.get(m.get('kieu'), m.get('kieu', '')),
               m.get('gia', ''), ', '.join(ten_buoi.get(b, b) for b in (m.get('buoi') or [])),
               m.get('mo_ta', ''), m.get('goi_y_cho', ''), '']
              for i, m in enumerate(mon['mon'])])

    to_sheet(wb, 'Lý do ăn', 'Câu hiện ra kèm món vừa quay được.',
             [('Mã (đừng sửa)', 14), ('Câu', 74), (GHI_CHU, 30)],
             [['ly_do' + str(i + 1), x, ''] for i, x in enumerate(mon['ly_do'])], loc=False)

    # --- tu tieng Nghe
    s = open(os.path.join(GOC, 'game.js'), encoding='utf-8').read()
    khoi = s[s.index('var TU = ['):]
    khoi = khoi[:khoi.index('\n  ];')]
    tu = re.findall(r"\['((?:[^'\\]|\\.)*)',\s*'((?:[^'\\]|\\.)*)',\s*'((?:[^'\\]|\\.)*)'\]", khoi)
    to_sheet(wb, 'Từ tiếng Nghệ',
             'Trò "Giọng Nghệ tốc độ" trong trang Game. Mỗi lượt lấy 10 từ ngẫu nhiên trong bảng này.',
             [('Mã (đừng sửa)', 12), ('Từ tiếng Nghệ', 20), ('Nghĩa', 24), ('Câu ví dụ', 50), (GHI_CHU, 26)],
             [['tu' + str(i + 1), a.replace("\\'", "'"), b.replace("\\'", "'"), c.replace("\\'", "'"), '']
              for i, (a, b, c) in enumerate(tu)])

    trang_huong_dan(wb, tv, mon, [ws.title for ws in wb.worksheets], ban_do_chay)
    wb.active = 0
    wb.save(tuy.ra)
    so_dong = sum(ws.max_row - 2 for ws in wb.worksheets if ws.title != 'Hướng dẫn')
    print('Da ghi:', tuy.ra)
    print('So sheet:', len(wb.worksheets), '| so dong noi dung:', so_dong)


if __name__ == '__main__':
    sys.exit(main())
